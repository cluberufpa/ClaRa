"""Prepara uma cópia enxuta; não altera a fonte nem integra o pacote."""
from pathlib import Path
import json,hashlib
from openpyxl import load_workbook,Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

p=Path(__file__).resolve().parent
arquivos=json.loads((p/'arquivos_dryad.json').read_text(encoding='utf-8-sig'))['_embedded']['stash:files']
fontes={}
for nome in ('Cowichan_Data.xlsx','README.md'):
    meta=next(a for a in arquivos if a['path']==nome)
    h=hashlib.sha256((p/nome).read_bytes()).hexdigest()
    assert h==meta['digest'],nome
    fontes[nome]={'bytes':(p/nome).stat().st_size,'sha256':h,'origem':'Arquivo baixado pelo professor; hash conferido com Dryad versão 446982'}
(p/'fontes.json').write_text(json.dumps(fontes,ensure_ascii=False,indent=2),encoding='utf-8')
w=load_workbook(p/'Cowichan_Data.xlsx',read_only=True,data_only=True)
rows=list(w['Dataset'].values); headers=rows[0]; rows=rows[1:]
assert len(rows)==246 and len(headers)==34
campos=[('ano','Year','Ano da coleta: 2022 ou 2023.'),
('local','Site name','Um dos cinco locais; nomes geográficos originais preservados.'),
('sublocal','Sub-site','Código original do ponto de coleta, sem expansão não documentada.'),
('ambiente','Waterbody type','Lake = lago; River = rio; Estuary = estuário.'),
('sexo','Sex','Female = fêmea; Male = macho, identificação gonadal na fonte.'),
('massa_fresca_g','Wet mass (g)','Massa fresca individual, em gramas.'),
('comprimento_total_mm','Total length (mm)','Comprimento total, em mm. Valor original mantido; ver pendência da linha 235.')]
idx=[headers.index(c[1]) for c in campos]
base=[]
for numero,r in enumerate(rows,1):
    vals=[r[i] for i in idx]
    vals[3]={'Lake':'lago','River':'rio','Estuary':'estuário'}[vals[3]]
    vals[4]={'Female':'fêmea','Male':'macho'}[vals[4]]
    assert all(v is not None and v!='NA' for v in vals)
    base.append([str(numero)]+vals)
out=Workbook(); s=out.active; s.title='LEIA-ME'
info=[('Campo','Descrição'),('Situação','Versão enxuta para revisão; ainda não substitui walleye_erie.'),
('Fonte','May et al. (2026), Dryad; https://doi.org/10.5061/dryad.6m905qg99'),
('Licença','CC0-1.0, confirmada nos metadados da versão 446982.'),
('Conteúdo','246 peixes × 8 colunas. Fonte integral de 34 colunas preservada ao lado.'),
('Uso principal','Amostragem simples e estratificada proporcional, usando esta base como população finita didática.'),
('Limite da interpretação','Não supor que a coleta original foi probabilística de toda a população natural.'),
('Identificação','id_registro corresponde à ordem original; linha da fonte Excel = id_registro + 1. Não é marca individual de campo.'),
('Seleção','Retidas identificação, contexto da coleta, sexo, massa e comprimento. Sem filtrar peixes ou imputar valores.'),
('Colunas fora do recorte','Temperatura, massa seca, placas/ossos, fósforo, excreção, digestão e morfometria detalhada; não necessárias ao objetivo de amostragem.'),
('Revisão de comprimento','Registro 234 / linha 235 da fonte: comprimento total 51,09 mm e padrão 55 mm. Não é possível decidir qual está errado. Usar massa nos exercícios iniciais.'),
('Data','11/10/2026')]
for r in info:s.append(r)
s=out.create_sheet('Dados para revisar'); s.append(['id_registro']+[c[0] for c in campos])
for r in base:s.append(r)
s.auto_filter.ref=s.dimensions
s=out.create_sheet('Dicionario'); s.append(['Coluna','Coluna original','Significado'])
s.append(['id_registro','Derivado da ordem da fonte','Texto: 1 a 246; preserva a ligação com a linha original.'])
for c in campos:s.append(c)
s=out.create_sheet('Pendencias'); s.append(['Item','Registro','Descrição'])
s.append(['Comprimento','234','Total 51,09 mm; padrão 55 mm na fonte. Confirmar com os autores; não corrigido.'])
s.append(['Substituição','Conjunto','Avaliação aprovada; troca definitiva por walleye_erie ainda não realizada.'])
s.append(['Amostragem','Conjunto','Usar massa como resposta inicial; selecionar estratos antes do sorteio.'])
for s in out:
    s.freeze_panes='A2'
    for cell in s[1]:cell.fill=PatternFill('solid',fgColor='0F3B5F');cell.font=Font(name='Arial',color='FFFFFF',bold=True)
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name='Arial',size=11)
            c.alignment=Alignment(vertical='top',wrap_text=s.title!='Dados para revisar')
            if c.row%2==0:c.fill=PatternFill('solid',fgColor='E7EFEA')
    for col in range(1,s.max_column+1):s.column_dimensions[get_column_letter(col)].width=24 if s.title=='Dados para revisar' else (27 if col==1 else 85)
    if s.title!='Dados para revisar':
        for row in range(2,s.max_row+1):s.row_dimensions[row].height=48
arquivo=p/'cowichan_para_revisao.xlsx';out.save(arquivo);w.close()
v=load_workbook(arquivo,read_only=True,data_only=True)
assert list(v['Dados para revisar'].values)[1:]==[tuple(r) for r in base]
assert v['Dados para revisar'].max_row==247 and v['Dados para revisar'].max_column==8
assert not any(c.data_type=='e' for sh in v for row in sh for c in row)
v.close()
print('Revisão conferida: 246 x 8; hashes originais corretos; valores e ordem preservados.')
